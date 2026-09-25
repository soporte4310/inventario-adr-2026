import uuid

import openpyxl
import pandas as pd
from openpyxl.styles import Font, PatternFill

from django.contrib import messages
from django.core.exceptions import ValidationError
from django.core.mail import EmailMultiAlternatives
from django.conf import settings
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect
from django.template.loader import render_to_string
from django.urls import reverse, reverse_lazy
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from django.views.generic import CreateView, ListView, TemplateView, View

from accounts.mixins import GroupRequiredMixin
from accounts.views import CustomLoginView
from common.utils import procesar_imagen_en_memoria

from .forms import (
    AgregarProyectorForm, EnroqueActivoForm, EnroqueImpresoraForm,
    ImpresoraForm, NovedadRevisionForm,
)
from .mixins import MantencionLoginRequiredMixin
from .models import CicloRevision, EquipoMantenible, EvidenciaRevision, Impresora, RegistroAuditoria, RevisionEquipo
from .utils import GRUPOS_ADMIN, GRUPOS_REVISION, crear_nuevo_ciclo, limite_alcanzado, registrar_auditoria


class MantencionLoginView(CustomLoginView):
    """
    Login 'propio' de Mantención: usa exactamente la misma autenticación
    que el portal de inventario (CustomAuthenticationForm: mismas
    credenciales, mismo bloqueo por intentos fallidos, mismo aviso de
    contraseña por cambiar), pero con plantilla propia y destino final
    distinto, para que quienes entren por esta puerta lleguen a
    /mantencion/ y no al dashboard de inventario.
    """
    template_name = 'mantencion/pages/login.html'
    next_page = 'mantencion_home'


def _url_revision(tipo):
    """Nombre de la url de 'Revisión mensual' que corresponde a un tipo de equipo."""
    return 'mantencion_revision_proyectores' if tipo == EquipoMantenible.Tipo.PROYECTOR else 'mantencion_revision_impresoras'


class TipoEquipoMixin:
    """Cada vista concreta fija 'tipo' para reutilizar la misma lógica en Proyectores/Impresoras."""
    tipo = None

    def get_tipo_label(self):
        return 'Proyectores' if self.tipo == EquipoMantenible.Tipo.PROYECTOR else 'Impresoras'

    def get_tipo_label_singular(self):
        return 'proyector' if self.tipo == EquipoMantenible.Tipo.PROYECTOR else 'impresora'


# ---------------------------------------------------------------------------
# 1. Home
# ---------------------------------------------------------------------------

class HomeMantencionView(MantencionLoginRequiredMixin, GroupRequiredMixin, TemplateView):
    group_required = GRUPOS_REVISION
    template_name = 'mantencion/pages/home.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)

        # Cada tipo tiene su propio ciclo más reciente (frecuencias
        # distintas), así que el mini-dashboard se arma por separado.
        for tipo, clave in [(EquipoMantenible.Tipo.PROYECTOR, 'proyectores'), (EquipoMantenible.Tipo.IMPRESORA, 'impresoras')]:
            ciclo = CicloRevision.objects.filter(tipo=tipo).first()
            qs = RevisionEquipo.objects.filter(ciclo=ciclo) if ciclo else RevisionEquipo.objects.none()
            ultima = qs.filter(fecha_revision__isnull=False).order_by('-fecha_revision').first()
            context[f'{clave}_ciclo'] = ciclo
            context[f'{clave}_pendientes'] = qs.filter(estado=RevisionEquipo.Estado.PENDIENTE).count()
            context[f'{clave}_ok'] = qs.filter(estado=RevisionEquipo.Estado.OK).count()
            context[f'{clave}_novedad'] = qs.filter(estado=RevisionEquipo.Estado.NOVEDAD).count()
            context[f'{clave}_ultima_revision'] = ultima.fecha_revision if ultima else None

        context['ultimas_novedades'] = RevisionEquipo.objects.filter(
            estado=RevisionEquipo.Estado.NOVEDAD
        ).select_related('equipo__activo', 'equipo__impresora', 'revisado_por').order_by('-fecha_revision')[:5]
        return context


# ---------------------------------------------------------------------------
# 2 y 3. Listas de equipos (datos detallados, sin acciones de revisión)
# ---------------------------------------------------------------------------

class ListaEquiposView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, ListView):
    group_required = GRUPOS_REVISION
    model = EquipoMantenible
    template_name = 'mantencion/pages/lista_equipos.html'
    context_object_name = 'equipos'

    def get_queryset(self):
        qs = EquipoMantenible.objects.filter(tipo=self.tipo, activo_en_revision=True).select_related(
            'activo__catalogo__marca', 'activo__catalogo__categoria', 'activo__estado',
            'activo__ubicacion__piso__edificio', 'impresora__ubicacion__piso__edificio',
        )
        # El orden "por edificio" sale de un campo real de BD, y proyectores
        # e impresoras lo guardan en modelos distintos (Activo vs Impresora),
        # así que el nombre del campo a ordenar depende del tipo.
        if self.tipo == EquipoMantenible.Tipo.PROYECTOR:
            return qs.order_by('activo__ubicacion__piso__edificio__nombre', 'activo__ubicacion__nombre')
        return qs.order_by('impresora__ubicacion__piso__edificio__nombre', 'impresora__ubicacion__nombre')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['tipo_label'] = self.get_tipo_label()
        context['tipo_label_singular'] = self.get_tipo_label_singular()
        context['tipo_code'] = self.tipo
        context['enroque_form'] = EnroqueActivoForm() if self.tipo == EquipoMantenible.Tipo.PROYECTOR else EnroqueImpresoraForm()
        if self.tipo == EquipoMantenible.Tipo.IMPRESORA:
            context['impresora_form'] = ImpresoraForm()
        # Modal "Ver equipos de baja": mismo campo que ya controla por qué un
        # equipo deja de aparecer en esta lista (pausado por ADR, o impresora
        # marcada De Baja, que también pausa su EquipoMantenible).
        context['equipos_de_baja'] = EquipoMantenible.objects.filter(
            tipo=self.tipo, activo_en_revision=False
        ).select_related(
            'activo__catalogo__marca', 'activo__catalogo__categoria', 'activo__estado',
            'activo__ubicacion__piso__edificio', 'impresora__ubicacion__piso__edificio',
        )
        return context


class ListaProyectoresView(ListaEquiposView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class ListaImpresorasView(ListaEquiposView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


# ---------------------------------------------------------------------------
# 4. Revisión mensual
# ---------------------------------------------------------------------------

def _ultimo_ciclo(tipo):
    """El ciclo más reciente de un tipo: es el único editable, el resto es historial de sólo lectura."""
    return CicloRevision.objects.filter(tipo=tipo).first()


class RevisionMensualView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, ListView):
    group_required = GRUPOS_REVISION
    template_name = 'mantencion/pages/revision_mensual.html'
    context_object_name = 'revisiones'

    def get_queryset(self):
        ultimo = _ultimo_ciclo(self.tipo)

        # ?ciclo=<id> deja ver un ciclo anterior (historial, sólo lectura);
        # sin ese parámetro, mostramos el más reciente (el editable).
        ciclo_id = self.request.GET.get('ciclo')
        if ciclo_id:
            self.ciclo = get_object_or_404(CicloRevision, pk=ciclo_id, tipo=self.tipo)
        else:
            self.ciclo = ultimo

        self.es_editable = bool(self.ciclo) and bool(ultimo) and self.ciclo.pk == ultimo.pk

        if not self.ciclo:
            return RevisionEquipo.objects.none()

        # 'ubicacion_en_revision' es un campo directo de RevisionEquipo (no
        # pasa por Activo/Impresora), así que este orden sirve para ambos
        # tipos sin necesidad de branch.
        return RevisionEquipo.objects.filter(ciclo=self.ciclo).select_related(
            'equipo__activo__catalogo__marca', 'equipo__activo__catalogo__categoria',
            'equipo__impresora', 'ubicacion_en_revision__piso__edificio', 'revisado_por'
        ).order_by('ubicacion_en_revision__piso__edificio__nombre', 'ubicacion_en_revision__nombre')

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['ciclo'] = self.ciclo
        context['ciclos'] = CicloRevision.objects.filter(tipo=self.tipo)
        context['es_editable'] = self.es_editable
        context['tipo_label'] = self.get_tipo_label()
        context['tipo_label_singular'] = self.get_tipo_label_singular()
        context['tipo_code'] = self.tipo
        context['novedad_form'] = NovedadRevisionForm()
        # Para la leyenda "Faltan por revisar X ..." y el resaltado amarillo
        # de las tarjetas que el practicante todavía no revisó en terreno.
        context['pendientes_count'] = self.object_list.filter(estado=RevisionEquipo.Estado.PENDIENTE).count()
        return context


class RevisionProyectoresView(RevisionMensualView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class RevisionImpresorasView(RevisionMensualView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


class NuevoCicloView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, View):
    """
    Botón "Nuevo ciclo/revisión": abre explícitamente una ronda nueva (con
    la fecha de hoy) para un tipo de equipo. Al abrirse, el ciclo anterior
    de ese mismo tipo pasa automáticamente a ser historial de sólo lectura
    (deja de ser "el más reciente").
    """
    group_required = GRUPOS_REVISION

    def post(self, request, *args, **kwargs):
        ciclo, cantidad = crear_nuevo_ciclo(self.tipo, usuario=request.user)
        messages.success(request, f"{ciclo} iniciado con {cantidad} equipos por revisar.")
        return redirect(_url_revision(self.tipo))


class NuevoCicloProyectoresView(NuevoCicloView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class NuevoCicloImpresorasView(NuevoCicloView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


class MarcarOkView(MantencionLoginRequiredMixin, GroupRequiredMixin, View):
    """Botón verde 'OK' de la revisión mensual."""
    group_required = GRUPOS_REVISION

    def post(self, request, pk):
        revision = get_object_or_404(
            RevisionEquipo.objects.select_related('equipo__activo', 'equipo__impresora', 'ciclo'), pk=pk
        )

        if limite_alcanzado(request.user, 'marcar_ok'):
            messages.error(request, "Demasiadas solicitudes seguidas. Espera un momento e intenta de nuevo.")
            return redirect(_url_revision(revision.equipo.tipo))

        ultimo = _ultimo_ciclo(revision.equipo.tipo)
        if not ultimo or revision.ciclo_id != ultimo.pk:
            messages.error(request, "Este ciclo ya quedó cerrado (es historial de sólo lectura); no se puede modificar.")
            return redirect(_url_revision(revision.equipo.tipo))

        revision.estado = RevisionEquipo.Estado.OK
        revision.revisado_por = request.user
        revision.fecha_revision = timezone.now()
        revision.save()

        registrar_auditoria(
            usuario=request.user, accion=RegistroAuditoria.Accion.REVISION_OK,
            equipo_descripcion=f"{revision.equipo.marca_modelo} (S/N {revision.equipo.numero_serie or '—'})",
            tipo_equipo=revision.equipo.tipo,
        )

        messages.success(request, f"{revision.equipo.equipo_real} marcado como OK.")
        # El ancla deja al navegador en la misma tarjeta que se acaba de
        # marcar, en vez de saltar arriba de todo con cada click. redirect()
        # necesita la ruta ya resuelta (reverse()): un nombre de url con '#'
        # pegado no se puede resolver como nombre de vista.
        url = reverse(_url_revision(revision.equipo.tipo))
        return redirect(f"{url}#revision-{revision.pk}")


class RegistrarNovedadView(MantencionLoginRequiredMixin, GroupRequiredMixin, View):
    """Botón 'Comentarios y evidencia' de la revisión mensual."""
    group_required = GRUPOS_REVISION

    def post(self, request, pk):
        revision = get_object_or_404(
            RevisionEquipo.objects.select_related('equipo__activo', 'equipo__impresora', 'ciclo'), pk=pk
        )

        if limite_alcanzado(request.user, 'novedad'):
            messages.error(request, "Demasiadas solicitudes seguidas. Espera un momento e intenta de nuevo.")
            return redirect(_url_revision(revision.equipo.tipo))

        ultimo = _ultimo_ciclo(revision.equipo.tipo)
        if not ultimo or revision.ciclo_id != ultimo.pk:
            messages.error(request, "Este ciclo ya quedó cerrado (es historial de sólo lectura); no se puede modificar.")
            return redirect(_url_revision(revision.equipo.tipo))

        form = NovedadRevisionForm(request.POST, request.FILES)
        if not form.is_valid():
            for error in form.errors.values():
                messages.error(request, error.as_text())
            return redirect(_url_revision(revision.equipo.tipo))

        archivo = form.cleaned_data['archivo']

        # Procesamos la evidencia ANTES de tocar la revisión: si la foto
        # viene corrupta o demasiado pesada, no queremos dejar la revisión
        # marcada como NOVEDAD sin ninguna evidencia adjunta.
        if form.tipo_archivo == EvidenciaRevision.TipoArchivo.FOTO:
            try:
                archivo_final = procesar_imagen_en_memoria(
                    image_field=archivo, max_dimensions=(1600, 1600),
                    new_filename=f"evidencia_{uuid.uuid4()}.jpg"
                )
            except ValidationError as error:
                messages.error(request, f"No se pudo procesar la foto: {'; '.join(error.messages)}")
                return redirect(_url_revision(revision.equipo.tipo))
        else:
            archivo_final = archivo

        sigue_operativo = form.cleaned_data['sigue_operativo']

        revision.estado = RevisionEquipo.Estado.NOVEDAD
        revision.comentario = form.cleaned_data['comentario']
        revision.sigue_operativo = sigue_operativo
        revision.revisado_por = request.user
        revision.fecha_revision = timezone.now()
        revision.save()

        EvidenciaRevision.objects.create(
            revision=revision, archivo=archivo_final,
            tipo_archivo=form.tipo_archivo, subido_por=request.user
        )

        etiqueta_operatividad = 'sigue operativo' if sigue_operativo else 'FUERA DE SERVICIO'
        registrar_auditoria(
            usuario=request.user, accion=RegistroAuditoria.Accion.REVISION_NOVEDAD,
            equipo_descripcion=f"{revision.equipo.marca_modelo} (S/N {revision.equipo.numero_serie or '—'})",
            tipo_equipo=revision.equipo.tipo,
            detalle=f"{form.cleaned_data['comentario']} ({etiqueta_operatividad})",
        )

        messages.warning(request, f"Novedad registrada para {revision.equipo.equipo_real}.")
        url = reverse(_url_revision(revision.equipo.tipo))
        return redirect(f"{url}#revision-{revision.pk}")


class EnroqueView(MantencionLoginRequiredMixin, GroupRequiredMixin, View):
    """
    El 'enroque': reasigna la sala/oficina de un equipo. Un proyector mueve
    su Activo real (auditado automáticamente por 'inventario'); una
    impresora mueve directamente su propio registro Impresora.

    Sólo ADR: mover un equipo de sala requiere planificación previa (no es
    algo que un practicante deba decidir sobre la marcha en terreno).
    """
    group_required = GRUPOS_ADMIN

    def post(self, request, equipo_id):
        equipo = get_object_or_404(
            EquipoMantenible.objects.select_related('activo__ubicacion', 'impresora__ubicacion'), pk=equipo_id
        )
        # Se captura ANTES de validar el form: ModelForm.is_valid() ya deja
        # el 'instance' (que es equipo.activo/impresora) con el valor NUEVO,
        # así que después de eso ya sería tarde para saber cuál era el viejo.
        ubicacion_anterior = str(equipo.ubicacion) if equipo.ubicacion else 'Sin ubicación'

        if equipo.tipo == EquipoMantenible.Tipo.PROYECTOR:
            form = EnroqueActivoForm(request.POST, instance=equipo.activo)
        else:
            form = EnroqueImpresoraForm(request.POST, instance=equipo.impresora)

        if form.is_valid():
            form.save()
            registrar_auditoria(
                usuario=request.user, accion=RegistroAuditoria.Accion.ENROQUE,
                equipo_descripcion=f"{equipo.marca_modelo} (S/N {equipo.numero_serie or '—'})",
                tipo_equipo=equipo.tipo,
                detalle=f"{ubicacion_anterior} → {equipo.ubicacion}",
            )
            messages.success(request, f"{equipo.equipo_real} fue reubicado correctamente.")
        else:
            messages.error(request, "No se pudo actualizar la ubicación. Selecciona una sala válida.")
        return redirect(_url_revision(equipo.tipo))


# ---------------------------------------------------------------------------
# Sólo ADR: alta/baja de impresoras (modal en "Lista de impresoras")
# ---------------------------------------------------------------------------

class ImpresoraAgregarView(MantencionLoginRequiredMixin, GroupRequiredMixin, View):
    """
    Modal "Agregar impresora": crea una Impresora nueva (no viene de
    'inventario', así que se cargan sus datos directamente) y la inscribe
    de inmediato en el circuito de mantención.
    """
    group_required = GRUPOS_ADMIN

    def post(self, request, *args, **kwargs):
        form = ImpresoraForm(request.POST)
        if not form.is_valid():
            for error in form.errors.values():
                messages.error(request, error.as_text())
            return redirect('mantencion_lista_impresoras')

        impresora = form.save(commit=False)
        impresora.creado_por = request.user
        impresora.save()

        EquipoMantenible.objects.create(
            tipo=EquipoMantenible.Tipo.IMPRESORA, impresora=impresora, agregado_por=request.user
        )

        registrar_auditoria(
            usuario=request.user, accion=RegistroAuditoria.Accion.IMPRESORA_AGREGADA,
            equipo_descripcion=str(impresora), tipo_equipo=EquipoMantenible.Tipo.IMPRESORA,
        )

        messages.success(request, f"{impresora} agregada al circuito de mantención.")
        return redirect('mantencion_lista_impresoras')


# ---------------------------------------------------------------------------
# Reportes (botón "Enviar reporte")
# ---------------------------------------------------------------------------

def _generar_excel_response(df, nombre_archivo, nombre_hoja='Datos'):
    """Arma la respuesta .xlsx con el mismo estilo de cabecera en todos los reportes de mantención."""
    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{nombre_archivo}.xlsx"'

    with pd.ExcelWriter(response, engine='openpyxl') as writer:
        df.to_excel(writer, index=False, sheet_name=nombre_hoja)
        worksheet = writer.sheets[nombre_hoja]
        for cell in worksheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="C00000", end_color="C00000", fill_type="solid")
        for idx, col in enumerate(df.columns):
            worksheet.column_dimensions[openpyxl.utils.get_column_letter(idx + 1)].width = max(len(col) + 2, 15)

    return response


def _dataframe_lista(tipo):
    """Datos del listado de equipos (no de un ciclo de revisión): lo que se ve en 'Lista de X'."""
    equipos = EquipoMantenible.objects.filter(tipo=tipo, activo_en_revision=True).select_related(
        'activo__catalogo__marca', 'activo__estado', 'activo__ubicacion__piso__edificio',
        'impresora__ubicacion__piso__edificio',
    )

    data = []
    for equipo in equipos:
        ubicacion = equipo.ubicacion
        data.append({
            'EDIFICIO': ubicacion.piso.edificio.nombre if ubicacion else '',
            'SALA_OFICINA': ubicacion.nombre if ubicacion else '',
            'MARCA_MODELO': equipo.marca_modelo,
            'NUMERO_SERIE': equipo.numero_serie,
            'ESTADO': equipo.estado_texto,
        })
    return pd.DataFrame(data)


class ListaExcelView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, View):
    """Excel del listado completo (no de un ciclo mensual), botón 'Excel' en Lista de proyectores/impresoras."""
    group_required = GRUPOS_REVISION

    def get(self, request, *args, **kwargs):
        df = _dataframe_lista(self.tipo)
        nombre_tipo = self.get_tipo_label()
        return _generar_excel_response(df, f"Listado_{nombre_tipo}", nombre_hoja='Listado')


class ListaExcelProyectoresView(ListaExcelView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class ListaExcelImpresoresView(ListaExcelView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


def _dataframe_revision(ciclo, tipo):
    if not ciclo:
        return pd.DataFrame()

    revisiones = RevisionEquipo.objects.filter(ciclo=ciclo, equipo__tipo=tipo).select_related(
        'equipo__activo__catalogo__marca', 'equipo__impresora',
        'ubicacion_en_revision__piso__edificio', 'revisado_por'
    ).order_by('ubicacion_en_revision__piso__edificio__nombre', 'ubicacion_en_revision__nombre')

    data = []
    for revision in revisiones:
        data.append({
            'EDIFICIO': revision.ubicacion_en_revision.piso.edificio.nombre if revision.ubicacion_en_revision else '',
            'SALA_OFICINA': revision.ubicacion_en_revision.nombre if revision.ubicacion_en_revision else '',
            'MARCA_MODELO': revision.equipo.marca_modelo,
            'NUMERO_SERIE': revision.equipo.numero_serie,
            'ESTADO': revision.get_estado_display(),
            'COMENTARIO': revision.comentario,
            'REVISADO_POR': revision.revisado_por.get_full_name() or revision.revisado_por.username if revision.revisado_por else '',
            'FECHA_REVISION': revision.fecha_revision.strftime('%d/%m/%Y %H:%M') if revision.fecha_revision else '',
        })
    return pd.DataFrame(data)


def _quedan_pendientes(ciclo):
    """El reporte (Excel o correo) solo tiene sentido con el ciclo completo: todavía no refleja la realidad si faltan equipos por revisar."""
    return bool(ciclo) and RevisionEquipo.objects.filter(ciclo=ciclo, estado=RevisionEquipo.Estado.PENDIENTE).exists()


class ReporteExcelView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, View):
    group_required = GRUPOS_REVISION

    def get(self, request, *args, **kwargs):
        ciclo_id = request.GET.get('ciclo')
        if ciclo_id:
            ciclo = get_object_or_404(CicloRevision, pk=ciclo_id, tipo=self.tipo)
        else:
            ciclo = _ultimo_ciclo(self.tipo)

        if _quedan_pendientes(ciclo):
            messages.error(request, "Todavía quedan equipos pendientes de revisar en este ciclo. El reporte se habilita cuando todos queden en OK o con novedad.")
            return redirect(_url_revision(self.tipo))

        df = _dataframe_revision(ciclo, self.tipo)
        nombre_tipo = self.get_tipo_label()
        etiqueta_fecha = ciclo.fecha_inicio.strftime('%d_%m_%Y') if ciclo else timezone.localdate().strftime('%d_%m_%Y')
        return _generar_excel_response(df, f"Revision_{nombre_tipo}_{etiqueta_fecha}", nombre_hoja='Revision')


class ReporteExcelProyectoresView(ReporteExcelView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class ReporteExcelImpresoresView(ReporteExcelView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


class ReporteEmailView(MantencionLoginRequiredMixin, GroupRequiredMixin, TipoEquipoMixin, View):
    group_required = GRUPOS_REVISION

    def post(self, request, *args, **kwargs):
        ciclo_id = request.POST.get('ciclo')
        if ciclo_id:
            ciclo = get_object_or_404(CicloRevision, pk=ciclo_id, tipo=self.tipo)
        else:
            ciclo = _ultimo_ciclo(self.tipo)

        if _quedan_pendientes(ciclo):
            messages.error(request, "Todavía quedan equipos pendientes de revisar en este ciclo. El reporte se habilita cuando todos queden en OK o con novedad.")
            return redirect(_url_revision(self.tipo))

        revisiones = RevisionEquipo.objects.filter(ciclo=ciclo) if ciclo else RevisionEquipo.objects.none()
        nombre_tipo = self.get_tipo_label()

        context = {
            'tipo': nombre_tipo,
            'ciclo': ciclo,
            'total': revisiones.count(),
            'ok': revisiones.filter(estado=RevisionEquipo.Estado.OK).count(),
            'pendientes': revisiones.filter(estado=RevisionEquipo.Estado.PENDIENTE).count(),
            'novedades': revisiones.filter(estado=RevisionEquipo.Estado.NOVEDAD).select_related(
                'equipo__activo', 'equipo__impresora', 'ubicacion_en_revision__piso__edificio'
            ),
        }
        html_content = render_to_string('mantencion/emails/reporte_revision.html', context)

        email = EmailMultiAlternatives(
            subject=f"Revisión de {nombre_tipo} - {ciclo}" if ciclo else f"Revisión de {nombre_tipo}",
            body="Favor visualizar este correo en modo HTML.",
            from_email=settings.DEFAULT_FROM_EMAIL,
            to=settings.EMAIL_RECIPIENTS,
        )
        email.attach_alternative(html_content, "text/html")
        email.send()

        messages.success(request, f"Reporte de {nombre_tipo.lower()} enviado a {len(settings.EMAIL_RECIPIENTS)} destinatarios.")
        return redirect(_url_revision(self.tipo))


class ReporteEmailProyectoresView(ReporteEmailView):
    tipo = EquipoMantenible.Tipo.PROYECTOR


class ReporteEmailImpresoresView(ReporteEmailView):
    tipo = EquipoMantenible.Tipo.IMPRESORA


# ---------------------------------------------------------------------------
# Sólo ADR: roster de equipos en mantención
# ---------------------------------------------------------------------------

class RosterEquiposView(MantencionLoginRequiredMixin, GroupRequiredMixin, ListView):
    group_required = GRUPOS_ADMIN
    model = EquipoMantenible
    template_name = 'mantencion/pages/roster_lista.html'
    context_object_name = 'equipos'

    def get_queryset(self):
        return EquipoMantenible.objects.select_related(
            'activo__catalogo__marca', 'activo__catalogo__categoria', 'activo__ubicacion__piso__edificio',
            'impresora__ubicacion__piso__edificio',
        ).order_by('tipo', '-fecha_agregado')


class RosterAgregarEquipoView(MantencionLoginRequiredMixin, GroupRequiredMixin, CreateView):
    """
    Agregar un PROYECTOR (Activo ya inventariado) al roster. Las impresoras
    ya no pasan por acá: se agregan desde el modal de "Lista de
    impresoras" (ImpresoraAgregarView), porque no vienen de 'inventario'.
    """
    group_required = GRUPOS_ADMIN
    model = EquipoMantenible
    form_class = AgregarProyectorForm
    template_name = 'mantencion/pages/roster_form.html'
    success_url = reverse_lazy('mantencion_roster')

    def form_valid(self, form):
        form.instance.agregado_por = self.request.user
        response = super().form_valid(form)

        # Si el ADR confirmó/corrigió la sala al agregar el equipo,
        # actualizamos el Activo real (mismo mecanismo que el 'enroque':
        # queda auditado automáticamente por las señales de inventario).
        ubicacion = form.cleaned_data.get('ubicacion')
        if ubicacion:
            activo = form.instance.activo
            activo.ubicacion = ubicacion
            activo.save()

        registrar_auditoria(
            usuario=self.request.user, accion=RegistroAuditoria.Accion.EQUIPO_AGREGADO,
            equipo_descripcion=f"{self.object.marca_modelo} (S/N {self.object.numero_serie or '—'})",
            tipo_equipo=self.object.tipo,
        )

        messages.success(self.request, "Proyector agregado al circuito de mantención mensual.")
        return response


class RosterPausarEquipoView(MantencionLoginRequiredMixin, GroupRequiredMixin, View):
    """
    Pausa/reactiva un equipo del roster sin borrarlo nunca: sólo deja de
    generarle filas de revisión mientras esté pausado.
    """
    group_required = GRUPOS_ADMIN

    def post(self, request, pk):
        equipo = get_object_or_404(
            EquipoMantenible.objects.select_related('activo', 'impresora'), pk=pk
        )
        equipo.activo_en_revision = not equipo.activo_en_revision
        equipo.save(update_fields=['activo_en_revision'])
        estado = "reactivado" if equipo.activo_en_revision else "quitado de la lista"

        registrar_auditoria(
            usuario=request.user,
            accion=RegistroAuditoria.Accion.EQUIPO_REACTIVADO if equipo.activo_en_revision else RegistroAuditoria.Accion.EQUIPO_PAUSADO,
            equipo_descripcion=f"{equipo.marca_modelo} (S/N {equipo.numero_serie or '—'})",
            tipo_equipo=equipo.tipo,
        )

        messages.success(request, f"Equipo {estado} correctamente.")
        # Si el botón vino desde "Lista de proyectores/impresoras" (no desde
        # el roster de ADR), volvemos a esa misma página en vez de forzar
        # la redirección al roster. url_has_allowed_host_and_scheme evita
        # que 'next' se use como redirección abierta hacia otro sitio.
        siguiente = request.POST.get('next')
        if siguiente and url_has_allowed_host_and_scheme(siguiente, allowed_hosts={request.get_host()}):
            return redirect(siguiente)
        return redirect('mantencion_roster')


# ---------------------------------------------------------------------------
# Sólo ADR: bitácora de auditoría (solo lectura)
# ---------------------------------------------------------------------------

class AuditoriaMantencionView(MantencionLoginRequiredMixin, GroupRequiredMixin, ListView):
    """
    Bitácora de solo lectura: quién hizo qué y cuándo dentro del módulo.
    Nunca se edita ni se borra nada desde acá, solo se consulta.
    """
    group_required = GRUPOS_ADMIN
    model = RegistroAuditoria
    template_name = 'mantencion/pages/auditoria.html'
    context_object_name = 'registros'
    paginate_by = 50

    def get_queryset(self):
        qs = RegistroAuditoria.objects.select_related('usuario')
        tipo = self.request.GET.get('tipo')
        if tipo in EquipoMantenible.Tipo.values:
            qs = qs.filter(tipo_equipo=tipo)
        accion = self.request.GET.get('accion')
        if accion in RegistroAuditoria.Accion.values:
            qs = qs.filter(accion=accion)
        return qs

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['tipo_filtro'] = self.request.GET.get('tipo', '')
        context['accion_filtro'] = self.request.GET.get('accion', '')
        context['acciones'] = RegistroAuditoria.Accion.choices
        return context
